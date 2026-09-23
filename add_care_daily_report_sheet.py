"""给 {prefix}_care_orders_fy_{start}_{end}.xlsx 添加 sheet「年度养护日报明细」：
  关联「年度养护」（订单级） + 「养护日报表」字段（care 设备级 + care_task 技师）。

「养护日报表」= 后台页面 wwwroot/background/maintain/task_report.html（养护业务报表）
的导出，数据源是 CareController.GetReport + Models/CareReport.cs。本脚本按同一口径
从生产库还原，因此列含义与 24-25 雪季那份日报表逐字对应。

日报表列（19）：
  门店 / 流水号 / 类型 / 品牌 / 长度 / 角度 /
  修刃 / 打蜡 / 机打蜡 / 刮蜡 / 其它 / 备注 / 维修技师 / 附加费用 /
  接待 / 日期 / 时间 / 订单ID / 金额

  原日报表 18 列，本脚本按用户口径把「机打蜡」单列拆出（打蜡 = 热蜡 ∪ 打蜡）；
  「订单号」改名「订单ID」（= order.id），避免与年度养护的「订单号」（= order.code）重名。

字段来源（对照 CareReport.cs）：
  门店      ← 主 sheet 门店        （CareReport.shop         = order.shop）
  流水号    ← care.task_flow_code  （CareReport.task_flow_num）
  类型/品牌/长度 ← care.equipment / brand / scale
  角度      ← care.edge_degree     （CareReport.degree）
  修刃      ← care_task '修刃'     → staff.name（CareReport.edge）
  打蜡      ← care_task '打蜡'/'热蜡' → staff.name（CareReport.wax，两值合并）
  机打蜡    ← care_task '机打蜡'   → staff.name（GetReport 未覆盖，本表新增）
  刮蜡      ← care_task '刮蜡'     → staff.name（CareReport.unwax）
  其它      ← care.repair_memo     （CareReport.more）
  备注      ← care.memo            （CareReport.memo）
  维修技师  ← care_task '维修'     → staff.name（CareReport.jishi）
  附加费用  ← care.repair_charge   （CareReport.additional_fee）
  接待      ← 主 sheet 店员姓名    （CareReport.staff       = order.staff.name）
  日期/时间 ← 主 sheet 业务日期/业务时间（task_report.html 取 order.biz_date）
  订单ID    ← order.id
  金额      ← 主 sheet 支付合计    （CareReport.total_paid  = Order.paidAmount
                                     = SUM(order_payment.amount WHERE status='支付成功')）

  6 个订单级列（门店/接待/日期/时间/订单ID/金额）中的 5 个直接取自主 sheet 同行，
  不重新查库，保证与工作簿其余 sheet 口径一致、不产生漂移。

技师多人：同一 care 同一工序可能有多条 care_task（多人/多次执行），
  按 care_task.id 升序去重后用 ';' 连接。这是对 GetReport 的有意偏离——
  它用 FirstOrDefault() 只取第一个会丢人，而本工作簿的「年度养护明细」已是 ';' 连接。

形态：
  - 订单级列（「年度养护」原全部 N 列 + 日报表侧 6 个订单级列）一单多 care 纵向合并
  - care 级 13 列在右侧，每个 care 各一行
  - 一单无 care 行 → 保留单行，care 级 13 列全部留空（不写 '--'）
  - 有 care 但该工序没做 → 5 个技师列写 '--'（同 24-25 日报表写法）
  - 角度为 NULL → 留空（不复刻旧手工导出写 0 的习惯）
  - 多 care 订单（M ≥ 2）整行底色浅蓝 EAF2FB
  - 表头 1F4E78 蓝底白字粗体；freeze A2

幂等：「年度养护日报明细」存在则删重建；其他 sheet 不动。

⚠️ 不要重跑 export_care_orders_fy.py —— 它整本重建 xlsx，会抹掉本 sheet
   以及 支付明细 / 支付流水 / 年度养护明细。

用法：
  python3 add_care_daily_report_sheet.py \
      --xlsx wanlong_service_care_orders_fy_2025-05-01_2026-04-30.xlsx \
      --shop 万龙服务中心
  # Intel Mac（只有 Driver 13）需额外传 --conn 覆盖默认的 Driver 18
"""
import argparse
import os
import sys
from collections import defaultdict
from datetime import datetime, timedelta

import pyodbc
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

MAIN_SHEET = '年度养护'
DETAIL_SHEET = '年度养护日报明细'
HEADER_COLOR = '1F4E78'
MULTI_FILL = 'EAF2FB'

# 日报表 19 列，顺序照搬原表并在「打蜡」后插入「机打蜡」
DAILY_COLS = ['门店', '流水号', '类型', '品牌', '长度', '角度',
              '修刃', '打蜡', '机打蜡', '刮蜡', '其它', '备注', '维修技师', '附加费用',
              '接待', '日期', '时间', '订单ID', '金额']

# 订单级列：一单多 care 时纵向合并（值取自主 sheet 同行，订单ID 取自 DB）
ORDER_LEVEL_COLS = ['门店', '接待', '日期', '时间', '订单ID', '金额']
# 订单级列 → 主 sheet 表头名（订单ID 不在主 sheet 里，单独从 DB 带）
ORDER_COL_FROM_MAIN = {
    '门店': '门店',
    '接待': '店员姓名',
    '日期': '业务日期',
    '时间': '业务时间',
    '金额': '支付合计',
}
# 技师列：有 care 但该工序没做时写 '--'
STAFF_COLS = ['修刃', '打蜡', '机打蜡', '刮蜡', '维修技师']
NO_STAFF = '--'
MONEY_DETAIL = {'附加费用', '金额'}

# task_name → 列名（'打蜡' 与 '热蜡' 共用「打蜡」列，同 CareReport.wax 口径）
TASK_TO_COL = {
    '修刃': '修刃',
    '打蜡': '打蜡',
    '热蜡': '打蜡',
    '机打蜡': '机打蜡',
    '刮蜡': '刮蜡',
    '维修': '维修技师',
}

DEFAULT_CONN = ('DRIVER={ODBC Driver 18 for SQL Server};SERVER=tcp:100.28.143.19,1433;'
                'DATABASE=snowmeet_new;UID=claude;PWD=abcd123!@#;'
                'Encrypt=yes;TrustServerCertificate=yes;Connection Timeout=30;')


def parse_args():
    p = argparse.ArgumentParser(
        description='给养护财年 xlsx 添加「年度养护日报明细」合并 sheet',
        formatter_class=argparse.RawDescriptionHelpFormatter, epilog=__doc__)
    p.add_argument('--xlsx', required=True, help='目标 xlsx 路径')
    p.add_argument('--shop', required=True, help='店铺名（DB order.shop）')
    p.add_argument('--start', default='2025-05-01', help='biz_date 起始（含），默认 2025-05-01')
    p.add_argument('--end', default='2026-04-30', help='biz_date 截止（含），默认 2026-04-30')
    p.add_argument('--conn', default=DEFAULT_CONN,
                   help='ODBC 连接串；Intel Mac 需换成 ODBC Driver 13')
    return p.parse_args()


def fetch_care_rows(conn_str, shop, start, end_excl):
    """返回 {order_code: [care_row_dict, ...]}, [care_id...]

    过滤条件与 add_care_detail_merged_sheet.py 逐字相同，保证两个明细 sheet 行数一致。
    """
    cn = pyodbc.connect(conn_str)
    cur = cn.cursor()
    cur.execute("""SELECT o.code, c.id, o.id, c.task_flow_code, c.equipment, c.brand, c.scale,
                          c.edge_degree, c.repair_memo, c.memo, c.repair_charge
                   FROM care c JOIN [order] o ON o.id = c.order_id
                   WHERE o.shop = ? AND o.[type] = N'养护'
                     AND o.biz_date >= ? AND o.biz_date < ?
                     AND o.valid = 1 AND o.code IS NOT NULL AND LTRIM(RTRIM(o.code)) <> N''
                     AND c.valid = 1
                   ORDER BY o.id, c.id""", shop, start, end_excl)
    by_code = defaultdict(list)
    care_ids = []
    for r in cur.fetchall():
        code = r[0].strip() if r[0] else None
        if not code:
            continue
        care_ids.append(r[1])
        by_code[code].append({
            'care_id': r[1],
            '订单ID': r[2],
            '流水号': r[3],
            '类型': r[4], '品牌': r[5], '长度': r[6],
            '角度': r[7],
            '其它': r[8],
            '备注': r[9],
            '附加费用': r[10],
        })
    cn.close()
    return by_code, care_ids


def fetch_order_ids(conn_str, shop, start, end_excl):
    """返回 {order_code: order_id}

    care 查询只覆盖有 care 行的订单，但「订单ID」是订单级字段，无 care 的订单也该有。
    这里按与主 sheet 相同的订单级条件单独取一遍（不加 o.valid 过滤，因为主 sheet
    的「正/闭」列说明它可能包含 valid=0 的订单）。
    """
    cn = pyodbc.connect(conn_str)
    cur = cn.cursor()
    cur.execute("""SELECT o.code, MAX(o.id)
                   FROM [order] o
                   WHERE o.shop = ? AND o.[type] = N'养护'
                     AND o.biz_date >= ? AND o.biz_date < ?
                     AND o.code IS NOT NULL AND LTRIM(RTRIM(o.code)) <> N''
                   GROUP BY o.code""", shop, start, end_excl)
    return {r[0].strip(): r[1] for r in cur.fetchall() if r[0]}


def fetch_staff_by_care(conn_str, care_ids):
    """返回 {care_id: {staff_col: 'name;name;...'}}"""
    if not care_ids:
        return {}
    cn = pyodbc.connect(conn_str)
    cur = cn.cursor()
    # 用临时表传 ids（避免 IN list 长度限制）
    cur.execute('CREATE TABLE #ids (id INT PRIMARY KEY)')
    cur.fast_executemany = True
    cur.executemany('INSERT INTO #ids (id) VALUES (?)', [(i,) for i in care_ids])
    cur.execute("""SELECT ct.care_id, ct.task_name, s.name, ct.id
                   FROM care_task ct
                   JOIN #ids tmp ON tmp.id = ct.care_id
                   LEFT JOIN staff s ON s.id = ct.staff_id
                   WHERE ct.valid = 1 AND ct.staff_id IS NOT NULL AND s.name IS NOT NULL
                   ORDER BY ct.care_id, ct.id""")
    # care_id → col_name → list of names（按 ct.id 升序，去重保序）
    by_care = defaultdict(lambda: defaultdict(list))
    for r in cur.fetchall():
        care_id, task_name, name, _ = r
        col = TASK_TO_COL.get(task_name)
        if col is None:
            continue
        if name not in by_care[care_id][col]:
            by_care[care_id][col].append(name)
    cur.execute('DROP TABLE #ids')
    cn.close()
    return {cid: {col: ';'.join(names) for col, names in cols.items()}
            for cid, cols in by_care.items()}


def main():
    args = parse_args()
    xlsx = os.path.abspath(args.xlsx)
    end_excl = (datetime.strptime(args.end, '%Y-%m-%d') + timedelta(days=1)).strftime('%Y-%m-%d')

    if not os.path.exists(xlsx):
        raise SystemExit(f'xlsx 不存在: {xlsx}')

    print(f'读 SQL care 明细（{args.shop} / {args.start} ~ {end_excl}）...')
    by_code, care_ids = fetch_care_rows(args.conn, args.shop, args.start, end_excl)
    total_cares = sum(len(v) for v in by_code.values())
    multi_orders = sum(1 for v in by_code.values() if len(v) > 1)
    print(f'  覆盖订单: {len(by_code)}（含 {multi_orders} 单多 care），care 总数: {total_cares}')

    print(f'读 care_task 关联员工（{len(care_ids)} 个 care_id）...')
    staff_map = fetch_staff_by_care(args.conn, care_ids)
    print(f'  有员工的 care: {len(staff_map)}')

    print('读订单 id（含无 care 的订单）...')
    order_id_map = fetch_order_ids(args.conn, args.shop, args.start, end_excl)
    print(f'  订单号 → id 映射: {len(order_id_map)} 条')

    print(f'打开 {xlsx}')
    wb = load_workbook(xlsx)
    if MAIN_SHEET not in wb.sheetnames:
        raise SystemExit(f'缺主 sheet「{MAIN_SHEET}」: {wb.sheetnames}')

    # 幂等：删旧 DETAIL_SHEET
    if DETAIL_SHEET in wb.sheetnames:
        print(f'  「{DETAIL_SHEET}」已存在，删除重建（幂等）')
        del wb[DETAIL_SHEET]

    main_ws = wb[MAIN_SHEET]
    main_headers = [main_ws.cell(row=1, column=c).value for c in range(1, main_ws.max_column + 1)]
    n_order_cols = len(main_headers)
    print(f'  年度养护订单级列数 N = {n_order_cols}')

    main_rows = []
    for r in range(2, main_ws.max_row + 1):
        row = [main_ws.cell(row=r, column=c).value for c in range(1, n_order_cols + 1)]
        main_rows.append(row)

    if '订单号' not in main_headers:
        raise SystemExit('年度养护找不到「订单号」列')
    code_idx = main_headers.index('订单号')

    # 日报表侧 5 个订单级列在主 sheet 里的下标（'支付合计' 等重名列取第一个出现的）
    main_col_idx = {}
    for daily_col, main_col in ORDER_COL_FROM_MAIN.items():
        if main_col not in main_headers:
            raise SystemExit(f'年度养护找不到「{main_col}」列（日报表「{daily_col}」需要它）')
        main_col_idx[daily_col] = main_headers.index(main_col)

    # 订单级金额列（用于保留数字格式）
    money_order_cols = set()
    for ci, h in enumerate(main_headers, start=1):
        if h in ('支付合计', '退款合计', '订单结余', '维修费合计', '普通养护费合计',
                 '减免合计', '卡券减免合计', '养护直减合计', '应分账金额', '实分账金额',
                 '待分账金额', '支付总金额', '退款总金额') or (isinstance(h, str) and h.endswith('】金额')):
            money_order_cols.add(ci)

    # 新 sheet
    ws = wb.create_sheet(DETAIL_SHEET)
    header_font = Font(bold=True, color='FFFFFF', name='Calibri', size=11)
    header_fill = PatternFill('solid', fgColor=HEADER_COLOR)
    multi_fill = PatternFill('solid', fgColor=MULTI_FILL)
    center = Alignment(horizontal='center', vertical='center')

    all_headers = list(main_headers) + DAILY_COLS
    for ci, h in enumerate(all_headers, start=1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center

    # 需要纵向合并的列号：主 sheet 全部 N 列 + 日报表侧 6 个订单级列
    merge_col_idxs = list(range(1, n_order_cols + 1))
    for col_name in ORDER_LEVEL_COLS:
        merge_col_idxs.append(n_order_cols + 1 + DAILY_COLS.index(col_name))

    out_row = 2
    merges = []
    multi_row_ranges = []

    for main_row in main_rows:
        code = main_row[code_idx]
        key = code.strip() if isinstance(code, str) else None
        care_list = by_code.get(key, []) if key else []
        M = max(len(care_list), 1)

        # 订单级列（仅第一行写值，其余等同合并视觉）
        for ci, v in enumerate(main_row, start=1):
            ws.cell(row=out_row, column=ci, value=v)

        # 日报表侧 5 个订单级列：取主 sheet 同行的值
        for col_name, mi in main_col_idx.items():
            ci = n_order_cols + 1 + DAILY_COLS.index(col_name)
            c = ws.cell(row=out_row, column=ci, value=main_row[mi])
            if col_name in MONEY_DETAIL and c.value is not None:
                c.number_format = '0.00'
        # 订单ID 取自 DB（无 care 的订单走 order_id_map 兜底）
        oid = care_list[0]['订单ID'] if care_list else order_id_map.get(key)
        ws.cell(row=out_row, column=n_order_cols + 1 + DAILY_COLS.index('订单ID'), value=oid)

        # care 级列：每个 care 一行
        for k in range(M):
            r = out_row + k
            if care_list and k < len(care_list):
                care = care_list[k]
                staff_cols = staff_map.get(care['care_id'], {})
                for di, col_name in enumerate(DAILY_COLS):
                    if col_name in ORDER_LEVEL_COLS:
                        continue  # 订单级列已在块首写过
                    if col_name in STAFF_COLS:
                        dv = staff_cols.get(col_name, NO_STAFF)
                    else:
                        dv = care.get(col_name)
                    c = ws.cell(row=r, column=n_order_cols + 1 + di, value=dv)
                    if col_name in MONEY_DETAIL and dv is not None:
                        c.number_format = '0.00'
            # 无 care 时 care 级列全部留空（不写 '--'）

        if M > 1:
            for ci in merge_col_idxs:
                merges.append((ci, out_row, out_row + M - 1))
            multi_row_ranges.append((out_row, out_row + M - 1))

        out_row += M

    last_row = out_row - 1
    print(f'  写入数据行: {last_row - 1} 行（含合并展开）')

    for ci, start_r, end_r in merges:
        col_letter = get_column_letter(ci)
        ws.merge_cells(f'{col_letter}{start_r}:{col_letter}{end_r}')
    print(f'  合并区域: {len(merges)} 个（{len(multi_row_ranges)} 单 × {len(merge_col_idxs)} 列）')

    # 多 care 订单整行浅蓝
    for start_r, end_r in multi_row_ranges:
        for r in range(start_r, end_r + 1):
            for ci in range(1, ws.max_column + 1):
                cell = ws.cell(row=r, column=ci)
                if cell.fill.fill_type is None:
                    cell.fill = multi_fill

    # 订单级金额列格式
    for ci in money_order_cols:
        for r in range(2, last_row + 1):
            cell = ws.cell(row=r, column=ci)
            if cell.value is not None and isinstance(cell.value, (int, float)):
                cell.number_format = '0.00'

    ws.freeze_panes = 'A2'

    for ci in range(1, len(all_headers) + 1):
        col_letter = get_column_letter(ci)
        max_w = sum(2 if ord(ch) > 127 else 1 for ch in str(all_headers[ci - 1]))
        for r in range(2, min(last_row + 1, 200)):
            v = ws.cell(row=r, column=ci).value
            if v is None:
                continue
            s = v.strftime('%Y-%m-%d %H:%M:%S') if hasattr(v, 'strftime') else str(v)
            w = sum(2 if ord(ch) > 127 else 1 for ch in s)
            if w > max_w:
                max_w = w
        ws.column_dimensions[col_letter].width = min(max_w + 2, 36)

    print(f'保存 {xlsx}')
    wb.save(xlsx)
    print(f'  sheets: {load_workbook(xlsx, read_only=True).sheetnames}')
    print('完成')


if __name__ == '__main__':
    main()
